"""향수별 분위기 정도 점수(0·1·2)를 엑셀에 반영한다 (2026-10-13, D-52).

입력: 재평가 결과 JSON 파일들 [{"id","scores":{"calm":2,...},"why"}]  (분위기 0점은 생략)
동작:
  - perfumes 시트에 mood_scores("calm:2,clean:1"), moods_v1(이전 태그) 열을 더하고,
  - moods 열을 점수가 1 이상인 분위기(점수 높은 순)로 바꾼다.
  - 서로 반대인 분위기(시원한↔따뜻한)는 한쪽만 남긴다. (처음에는 시원한↔포근한도 막았으나 D-54에서 풀었다. 이미 매긴 점수에는 시원한+포근한이 함께 있는 향수가 없다.)
실행: python3 scripts/apply-mood-scores.py out0.json out1.json ...   (openpyxl 필요) -> 그다음 xlsx-to-json.py
"""
import json
import sys
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

XLSX = Path(__file__).resolve().parent.parent / "docs" / "perfumes-draft.xlsx"
MOODS = ["calm", "clean", "cool", "cozy", "fresh", "zesty", "warm", "sweet", "lovely", "elegant", "deep", "sensual", "soft"]
OPPOSITE = [("cool", "warm")]

results = {}
for f in sys.argv[1:]:
    for r in json.load(open(f, encoding="utf8")):
        results[r["id"]] = r

wb = load_workbook(XLSX)
ws = wb["perfumes"]
H = [c.value for c in ws[1]]


def col(name):
    if name not in H:
        H.append(name)
        ws.cell(1, len(H)).value = name
    return H.index(name) + 1


c_id, c_moods, c_int = col("id"), col("moods"), col("intensity")
c_scores, c_v1 = col("mood_scores"), col("moods_v1")
log = Counter()
missing = []
sizes = []
for row in range(2, ws.max_row + 1):
    pid = ws.cell(row, c_id).value
    if pid is None:
        continue
    old = [x.strip() for x in str(ws.cell(row, c_moods).value or "").split(",") if x.strip()]
    if not ws.cell(row, c_v1).value:
        ws.cell(row, c_v1).value = ", ".join(old)  # 처음 한 번만 이전 태그를 남긴다
    else:
        old = [x.strip() for x in str(ws.cell(row, c_v1).value).split(",") if x.strip()]
    r = results.get(pid)
    if not r:
        missing.append(pid)
        continue
    sc = {k: int(v) for k, v in r["scores"].items() if k in MOODS and int(v) in (1, 2)}
    for a, b in OPPOSITE:
        if sc.get(a) and sc.get(b):
            keep = a if (sc[a], a in old) >= (sc[b], b in old) else b
            drop = b if keep == a else a
            del sc[drop]
            log["반대 분위기 한쪽 제거"] += 1
    if ws.cell(row, c_int).value == "light" and sc.get("deep") == 2:
        sc["deep"] = 1
        log["가벼운 향의 묵직한 2→1"] += 1
    order = sorted(sc, key=lambda k: (-sc[k], MOODS.index(k)))
    ws.cell(row, c_moods).value = ", ".join(order)
    ws.cell(row, c_scores).value = ",".join(f"{k}:{sc[k]}" for k in order)
    sizes.append(len(order))
    log["이전 태그 중 빠진 것"] += len([m for m in old if m not in sc])
    log["새로 생긴 분위기"] += len([m for m in sc if m not in old])
wb.save(XLSX)
print("반영한 향수", len(sizes), "누락", missing[:5], len(missing))
print("분위기 개수(0 초과) 평균", round(sum(sizes) / max(1, len(sizes)), 2), "최소", min(sizes), "최대", max(sizes))
print(dict(log))
