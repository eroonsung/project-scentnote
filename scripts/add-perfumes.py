"""새 향수 행을 엑셀에 추가한다 (한 번에 여러 chunk 결과 파일을 합침).
입력: <WORK>/out*/result.json (perfumes, newNotes, brandChannels) + <WORK>/imgs/*/ 이미지
동작: 검증 → 엑셀 perfumes/notes 시트에 추가 → 이미지를 public/perfumes/에 복사
사용: python scripts/add-perfumes.py <WORK> [--write]   (--write 없으면 검증 결과만 출력)"""
import glob, json, os, shutil, sys
import openpyxl
from copy import copy

WORK = sys.argv[1]
WRITE = "--write" in sys.argv
XLSX = "docs/perfumes-draft.xlsx"
FAMILIES = "citrus fruity aqua green aromatic floral whiteFloral powdery musk spicy woody amber gourmand leather".split()
MOODS = "calm clean cool cozy fresh zesty warm sweet lovely elegant deep sensual soft".split()
AVOIDS = "tooSweet tooStrong headache heavyFloral smoky powdery spicy soapy".split()
CONFLICTS = [("cool", "warm"), ("cool", "cozy"), ("cool", "sensual"), ("cool", "deep"), ("clean", "sensual"), ("clean", "deep"),
             ("zesty", "deep"), ("zesty", "sensual"), ("fresh", "sensual"), ("soft", "sensual"), ("lovely", "sensual"), ("lovely", "deep")]
CONC = {"EDP", "EDT", "EDC", "Parfum", "Cologne", None}
split = lambda v: [x.strip() for x in str(v).split(",") if x and x.strip()] if v else []

wb = openpyxl.load_workbook(XLSX)
ws, nw = wb["perfumes"], wb["notes"]
head = [c.value for c in ws[1]]
col = {k: i + 1 for i, k in enumerate(head)}
existing = {ws.cell(r, col["id"]).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, col["id"]).value}
by_brand = {}
for r in existing.values():
    by_brand.setdefault(ws.cell(r, col["brand"]).value, r)
notes = {nw.cell(r, 1).value for r in range(2, nw.max_row + 1) if nw.cell(r, 1).value}

rows, new_notes, brand_ch, problems, dropped = [], [], {}, [], []
for f in sorted(glob.glob(os.path.join(WORK, "out*", "result.json"))):
    d = json.load(open(f, encoding="utf-8"))
    rows += d.get("perfumes", [])
    for n in d.get("newNotes", []):
        if n["en"] not in notes and n["en"] not in {x["en"] for x in new_notes}:
            new_notes.append(n)
    for b in d.get("brandChannels", []):
        brand_ch[b["brand"]] = b
    dropped += d.get("dropped", [])
all_notes = notes | {n["en"] for n in new_notes}

seen = set()
good = []
for p in rows:
    errs = []
    pid = p.get("id")
    if not pid: errs.append("id 없음")
    if pid in existing or pid in seen: errs.append("id 중복")
    if p.get("family") not in FAMILIES: errs.append(f"family {p.get('family')}")
    for s in split(p.get("sub_families")):
        if s not in FAMILIES: errs.append(f"sub {s}")
    ms = split(p.get("moods"))
    if not ms or any(m not in MOODS for m in ms): errs.append(f"moods {ms}")
    for a, b in CONFLICTS:
        if a in ms and b in ms: errs.append(f"충돌 {a}-{b}")
    for a in split(p.get("avoid_flags")):
        if a not in AVOIDS: errs.append(f"avoid {a}")
    if p.get("intensity") not in ("light", "medium", "strong"): errs.append("intensity")
    if p.get("polarizing") not in ("low", "mid", "high"): errs.append("polarizing")
    if p.get("concentration") not in CONC: errs.append(f"concentration {p.get('concentration')}")
    for k in ("top_notes", "middle_notes", "base_notes"):
        for n in split(p.get(k)):
            if n not in all_notes: errs.append(f"노트 사전에 없음: {n}")
    for k in ("summary", "drydown", "check_point", "name_ko", "brand_ko"):
        if not p.get(k): errs.append(f"{k} 비어 있음")
    if p.get("kr_krw") is not None and not isinstance(p["kr_krw"], int): errs.append("kr_krw 숫자 아님")
    if p.get("kr_krw") and not (p.get("kr_source_url") and p.get("kr_ml")): errs.append("가격 근거(url/ml) 없음")
    img = p.get("image_file")
    if img:
        found = glob.glob(os.path.join(WORK, "imgs", "*", img))
        if not found: errs.append(f"이미지 파일 없음 {img}")
        elif os.path.getsize(found[0]) > 1_600_000: errs.append(f"이미지 큼 {img}")
        elif not p.get("image_source_url"): errs.append("이미지 출처 없음")
    if errs:
        problems.append((pid, errs))
    else:
        seen.add(pid); good.append(p)

print(f"행 {len(rows)}개 중 통과 {len(good)}개, 문제 {len(problems)}개, dropped {len(dropped)}개, 새 노트 {len(new_notes)}개")
for pid, e in problems: print("  문제:", pid, e)
print("  가격 있음:", sum(1 for p in good if p.get("kr_krw")), "| 이미지 있음:", sum(1 for p in good if p.get("image_file")))
if not WRITE: sys.exit(0)

# 채널: 기존 브랜드는 같은 브랜드 행 값을 복사, 새 브랜드는 조사 결과(없으면 미확인)
for p in good:
    r = by_brand.get(p["brand"])
    if r:
        p["channels"] = ws.cell(r, col["channels"]).value
        p["channels_status"] = ws.cell(r, col["channels_status"]).value
    else:
        b = brand_ch.get(p["brand"], {})
        st = b.get("status", "미확인")
        p["channels"] = b.get("channels")
        p["channels_status"] = {"높음": "높음 (앱에 표시)", "보통": "보통 (앱에 표시)"}.get(st, f"{st} (앱에 표시 안 함)")

last = ws.max_row
tmpl = last
for i, p in enumerate(good):
    r = last + 1 + i
    for c in range(1, len(head) + 1):
        src, dst = ws.cell(tmpl, c), ws.cell(r, c)
        dst.font, dst.fill, dst.border, dst.alignment, dst.number_format = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment), src.number_format
        v = p.get(head[c - 1])
        dst.value = v if v != "" else None
new_last = last + len(good)

nlast = nw.max_row
for i, n in enumerate(new_notes):
    r = nlast + 1 + i
    for c in range(1, 6):
        src, dst = nw.cell(nlast, c), nw.cell(r, c)
        dst.font, dst.fill, dst.border, dst.alignment, dst.number_format = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment), src.number_format
    nw.cell(r, 1).value, nw.cell(r, 2).value, nw.cell(r, 3).value, nw.cell(r, 4).value = n["en"], n["ko"], n["feel"], n["family"]
for r in range(2, nw.max_row + 1):
    nw.cell(r, 5).value = f'=COUNTIF(perfumes!$G$2:$I${new_last},"*"&A{r}&"*")'

for dv in ws.data_validations.dataValidation:
    dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(str(dv.sqref).replace("200", "500"))
ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(head))}{new_last}"
o = wb["options"]
for row in o.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("=COUNTIF(perfumes!") and "$200" in c.value:
            c.value = c.value.replace("$200", "$500")

# 채널 시트: 새 브랜드 추가
cw = wb["channels"]
have = {cw.cell(r, 1).value for r in range(2, cw.max_row + 1)}
for b, v in brand_ch.items():
    if b in have: continue
    r = cw.max_row + 1
    for c in range(1, 7):
        s, d = cw.cell(r - 1, c), cw.cell(r, c)
        d.font, d.fill, d.border, d.alignment = copy(s.font), copy(s.fill), copy(s.border), copy(s.alignment)
    st = v.get("status", "미확인")
    cw.cell(r, 1).value, cw.cell(r, 2).value = b, st
    cw.cell(r, 3).value = "표시" if st in ("높음", "보통") else "표시 안 함"
    cw.cell(r, 4).value, cw.cell(r, 5).value = v.get("channels"), v.get("evidence_url")

os.makedirs("public/perfumes", exist_ok=True)
for p in good:
    if p.get("image_file"):
        src = glob.glob(os.path.join(WORK, "imgs", "*", p["image_file"]))[0]
        shutil.copy(src, os.path.join("public/perfumes", p["image_file"]))
wb.save(XLSX)
print("엑셀 저장 완료 (perfumes", new_last - 1, "행, 새 노트", len(new_notes), "개), 이미지 복사 완료")
