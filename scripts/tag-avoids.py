"""피하고 싶은 칩 태그(avoid_flags)를 향수별로 직접 정한 값으로 엑셀에 쓴다 (노트·계열·강도·일반적인 평판을 보고 사람이 검토하듯 판단한 AI 초안).
목록에 없는 향수는 태그 없음. 사용: python scripts/tag-avoids.py [--write]"""
import sys
import openpyxl

S, T, H, F, K, P, SP, SO = "tooSweet", "tooStrong", "headache", "heavyFloral", "smoky", "powdery", "spicy", "soapy"
FLAGS = {
    "calvin-klein-ck-one": [SO],
    "burberry-her": [S],
    "dior-sauvage-edt": [SP],
    "jean-paul-gaultier-le-male": [S, SP],
    "viktor-rolf-flowerbomb": [F, S, T, H],
    "gucci-bloom": [F],
    "diptyque-do-son": [F],
    "chanel-no5-edp": [P, F, SO],
    "prada-infusion-d-iris": [P],
    "dior-homme-intense": [P],
    "chloe-edp": [P],
    "byredo-blanche": [SO],
    "maison-margiela-replica-lazy-sunday-morning": [SO],
    "narciso-rodriguez-for-her-edp": [SO],
    "diptyque-fleur-de-peau": [P, SO],
    "hermes-terre-d-hermes-edt": [SP],
    "hermes-twilly-d-hermes": [SP],
    "le-labo-santal-33": [T, K],
    "tom-ford-oud-wood": [T, K, SP],
    "mfk-baccarat-rouge-540": [S, T],
    "guerlain-shalimar-edp": [T, P],
    "ysl-black-opium": [S, T],
    "lancome-la-vie-est-belle": [S, P],
    "kilian-angels-share": [S, T],
    "mugler-angel": [S, T, H],
    "maison-margiela-replica-by-the-fireplace": [K],
    "maison-margiela-replica-jazz-club": [K],
    "tom-ford-tobacco-vanille": [S, T, K, SP],
    "le-labo-rose-31": [T, SP, K],
    "le-labo-the-noir-29": [K],
    "kilian-love-dont-be-shy": [S, T],
    "kilian-black-phantom": [S, T],
    "creed-aventus": [K, T],
    "frederic-malle-portrait-of-a-lady": [T, F, SP, K],
    "frederic-malle-musc-ravageur": [T, SP],
    "frederic-malle-carnal-flower": [F, T, H],
    "frederic-malle-en-passant": [P],
    "frederic-malle-lipstick-rose": [P],
    "aesop-hwyl": [K],
    "parfums-de-marly-layton": [S],
    "mfk-aqua-universalis": [SO],
    "mfk-oud-satin-mood": [T, K],
    "mfk-grand-soir": [S, T],
    "initio-oud-for-greatness": [T, K, H],
    "diptyque-eau-duelle": [SP, K],
    "tom-ford-lost-cherry": [S, T],
    "tom-ford-tuscan-leather": [T, K],
    "tom-ford-black-orchid": [F, T, S, H],
    "tom-ford-ombre-leather": [T, K],
    "xerjoff-naxos": [S, T, SP, K],
    "parfums-de-marly-herod": [S, SP, K],
    "chanel-coromandel": [S, T],
    "versace-eros": [S],
    "paco-rabanne-1-million": [S, T, SP],
    "paco-rabanne-lady-million": [S],
    "hugo-boss-bottled": [SP],
    "carolina-herrera-good-girl": [S, T],
    "ysl-mon-paris": [S],
    "prada-candy": [S],
    "mugler-alien": [F, T, H],
}
# 분위기 태그 손질: 어울리지 않는 자동 추가 태그 제거
MOOD_DROP = {"versace-eros": ["calm"], "jean-paul-gaultier-le-male": ["calm"]}

wb = openpyxl.load_workbook("docs/perfumes-draft.xlsx")
ws = wb["perfumes"]
col = {c.value: i + 1 for i, c in enumerate(ws[1])}
ids = {ws.cell(r, col["id"]).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, col["id"]).value}
missing = [k for k in list(FLAGS) + list(MOOD_DROP) if k not in ids]
assert not missing, missing
print("플래그별 향수 수")
for f in (S, T, H, F, K, P, SP, SO): print(f"  {f}: {sum(f in v for v in FLAGS.values())}")
if "--write" in sys.argv:
    for pid, r in ids.items():
        ws.cell(r, col["avoid_flags"]).value = ", ".join(FLAGS.get(pid, [])) or None
    for pid, drop in MOOD_DROP.items():
        c = ws.cell(ids[pid], col["moods"])
        c.value = ", ".join(m for m in c.value.split(", ") if m not in drop)
    wb.save("docs/perfumes-draft.xlsx"); print("엑셀 저장 완료")
