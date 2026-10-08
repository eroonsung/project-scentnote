"""향수 분류 재검토 결과(out*.json)를 엑셀 perfumes 시트에 반영한다.
각 향수의 moods / avoid_flags / intensity / polarizing / family / sub_families를 검토 결과로 덮어쓴다.
사용: python scripts/apply-review.py <REVIEW_DIR> [--write]   (--write 없으면 검증·변경 요약만 출력)"""
import glob, json, os, sys
import openpyxl

DIR = sys.argv[1]
WRITE = "--write" in sys.argv
FAMILIES = set("citrus fruity aqua green aromatic floral whiteFloral powdery musk spicy woody amber gourmand leather".split())
MOODS = set("calm clean cool cozy fresh zesty warm sweet lovely elegant deep sensual soft".split())
AVOIDS = set("tooSweet tooStrong headache heavyFloral smoky powdery spicy soapy".split())
CONFLICTS = [("cool", "warm"), ("cool", "cozy"), ("cool", "sensual"), ("cool", "deep"), ("clean", "sensual"), ("clean", "deep"),
             ("zesty", "deep"), ("zesty", "sensual"), ("fresh", "sensual"), ("soft", "sensual"), ("lovely", "sensual"), ("lovely", "deep")]
split = lambda v: [x.strip() for x in str(v).split(",") if x and x.strip()] if v else []

wb = openpyxl.load_workbook("docs/perfumes-draft.xlsx")
ws = wb["perfumes"]
col = {c.value: i + 1 for i, c in enumerate(ws[1])}
rowof = {ws.cell(r, col["id"]).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, col["id"]).value}

reviews, problems = {}, []
for f in sorted(glob.glob(os.path.join(DIR, "out*.json"))):
    for d in json.load(open(f, encoding="utf-8")):
        reviews[d["id"]] = d
missing = [i for i in rowof if i not in reviews]
print(f"검토 결과 {len(reviews)}개 / 향수 {len(rowof)}개, 결과 없는 향수 {len(missing)}개")

changes = {"moods": 0, "avoid_flags": 0, "intensity": 0, "polarizing": 0, "family": 0, "sub_families": 0}
for pid, d in reviews.items():
    if pid not in rowof:
        problems.append((pid, "엑셀에 없는 id")); continue
    ms = split(d.get("moods")); av = split(d.get("avoid_flags")); subs = split(d.get("sub_families"))
    errs = []
    if not (3 <= len(ms) <= 5) or any(m not in MOODS for m in ms): errs.append(f"moods {ms}")
    for a, b in CONFLICTS:
        if a in ms and b in ms: errs.append(f"충돌 {a}-{b}")
    if any(a not in AVOIDS for a in av): errs.append(f"avoid {av}")
    if d.get("family") not in FAMILIES or any(s not in FAMILIES for s in subs): errs.append("family")
    if d.get("intensity") not in ("light", "medium", "strong") or d.get("polarizing") not in ("low", "mid", "high"): errs.append("intensity/polarizing")
    if errs:
        problems.append((pid, errs)); continue
    r = rowof[pid]
    new = {"moods": ", ".join(ms), "avoid_flags": ", ".join(av) or None, "intensity": d["intensity"], "polarizing": d["polarizing"],
           "family": d["family"], "sub_families": ", ".join(subs) or None}
    for k, v in new.items():
        cur = ws.cell(r, col[k]).value
        if (cur or None) != (v or None):
            changes[k] += 1
            if WRITE: ws.cell(r, col[k]).value = v
for p in problems: print("  문제:", p)
print("바뀐 칸 수:", changes)
if WRITE:
    wb.save("docs/perfumes-draft.xlsx"); print("엑셀 저장 완료")
