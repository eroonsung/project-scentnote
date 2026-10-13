"""올리브영 향수 Top 120 브랜드 목록(2026-10-13 확인)에 있는 브랜드의 채널에 '올리브영'을 반영한다.

온라인 브랜드 필터 목록이라 매장별 취급·시향 가능 여부는 보장하지 않는다.
실행: python3 scripts/apply-oliveyoung-brands.py   (openpyxl 필요) -> 그다음 xlsx-to-json.py
"""
from pathlib import Path
from openpyxl import load_workbook

XLSX = Path(__file__).resolve().parent.parent / "docs" / "perfumes-draft.xlsx"
# 올리브영 목록에 있는 우리 앱의 브랜드(영문명). 목록에 없는 브랜드는 건드리지 않는다.
OLIVEYOUNG = {"Narciso Rodriguez", "Davidoff", "Dolce & Gabbana", "Chloé", "Lancôme", "Marc Jacobs", "Montblanc",
              "Burberry", "Versace", "Issey Miyake", "Calvin Klein", "Elizabeth Arden"}
# 매장 검색어를 올리브영으로 돌리는 브랜드(지도에서 브랜드 매장이 안 잡히는 브랜드)
TO_OY_QUERY = {"Davidoff", "Calvin Klein"}
NOTE = "올리브영 향수 Top 120 브랜드 목록(2026-10-13, 온라인 브랜드 필터)에 있음. 매장별 취급·시향 가능 여부는 확인하지 못함."

wb = load_workbook(XLSX)
ws = wb["perfumes"]
H = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(H)}
updated = {}
for r in range(2, ws.max_row + 1):
    en = ws.cell(r, ix["brand"]).value
    if en not in OLIVEYOUNG:
        continue
    cur = [x.strip() for x in str(ws.cell(r, ix["channels"]).value or "").split(",") if x.strip() and "확인 필요" not in x]
    if not any("올리브영" in x for x in cur):  # "백화점 · 올리브영"처럼 묶인 항목도 인정
        cur.append("올리브영")
    ws.cell(r, ix["channels"]).value = ", ".join(cur)
    status = str(ws.cell(r, ix["channels_status"]).value or "")
    if not status.startswith(("높음", "보통")):
        ws.cell(r, ix["channels_status"]).value = "보통 (앱에 표시)"
    if en in TO_OY_QUERY:
        ws.cell(r, ix["store_query"]).value = "올리브영 향수"
    updated[en] = cur
wc = wb["channels"]
for r in range(2, wc.max_row + 1):
    en = wc.cell(r, 1).value
    if en in updated:
        wc.cell(r, 4).value = ", ".join(updated[en])
        if str(wc.cell(r, 2).value or "") not in ("높음", "보통"):
            wc.cell(r, 2).value = "보통"
            wc.cell(r, 3).value = "표시"
        note = str(wc.cell(r, 6).value or "")
        if NOTE not in note:
            wc.cell(r, 6).value = (note + " | " if note and note != "None" else "") + NOTE
wb.save(XLSX)
print(updated)
